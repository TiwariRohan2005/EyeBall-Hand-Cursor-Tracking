import unittest
import tempfile
import os

from library_app.database.db import DatabaseManager
from library_app.database.repository import LibraryRepository
from library_app.reports.report_service import ReportService

class TestReporting(unittest.TestCase):
    def setUp(self):
        self.temp_db = tempfile.NamedTemporaryFile(delete=False)
        self.temp_db.close()
        
        self.db = DatabaseManager(self.temp_db.name)
        self.db.init_db()
        self.repo = LibraryRepository(self.db)
        
        self.repo.add_user("OwnerUser", role="OWNER")
        self.repo.add_user("StaffUser", role="STAFF")
        
        self.repo.add_book({
            'book_code': 'B1', 
            'title': 'Test Report Book',
            'status': 'Available'
        })
        
        self.temp_reports_dir = tempfile.TemporaryDirectory()
        self.report_service = ReportService(self.repo)
        self.report_service.reports_dir = self.temp_reports_dir.name
            
    def tearDown(self):
        self.temp_reports_dir.cleanup()
        if os.path.exists(self.temp_db.name):
            try:
                os.unlink(self.temp_db.name)
            except:
                pass
                
    def test_owner_export_creates_file(self):
        # First export creates the persistent file
        path, is_new = self.report_service.generate_report("OWNER", "OwnerUser")
        self.assertTrue(os.path.exists(path))
        self.assertTrue(is_new)
        self.assertEqual(os.path.basename(path), "TrueEyeball_Library_Report.xlsx")
        
    def test_owner_export_updates_file(self):
        # Export twice to test update logic
        path1, is_new1 = self.report_service.generate_report("OWNER", "OwnerUser")
        self.assertTrue(is_new1)
        
        # Second time should update and not return is_new
        path2, is_new2 = self.report_service.generate_report("OWNER", "OwnerUser")
        self.assertEqual(path1, path2)
        self.assertFalse(is_new2)
        
        import openpyxl
        wb = openpyxl.load_workbook(path2)
        ws_txn = wb["Raw Transactions Log"]
        # There are currently 0 transactions in this setup, wait we should add one.
        
    def test_transactions_not_duplicated(self):
        # Create a transaction
        book = self.repo.get_book_by_code('B1')
        self.repo.issue_book(book['id'], "OwnerUser", "OwnerUser")
        
        path1, _ = self.report_service.generate_report("OWNER", "OwnerUser")
        import openpyxl
        wb1 = openpyxl.load_workbook(path1)
        # 1 header + 1 transaction row
        self.assertEqual(getattr(wb1["Raw Transactions Log"], 'max_row', 1), 2)
        
        # Export again
        path2, _ = self.report_service.generate_report("OWNER", "OwnerUser")
        wb2 = openpyxl.load_workbook(path2)
        # Should still be 2, not 3.
        self.assertEqual(getattr(wb2["Raw Transactions Log"], 'max_row', 1), 2)

    def test_file_in_use_error(self):
        # Create file first
        path, _ = self.report_service.generate_report("OWNER", "OwnerUser")
        
        # Mock PermissionError by manually holding the file handle exclusively via os.open in Windows? 
        # Easier to mock ExcelExporter.generate_workbook since cross-platform file locking in tests is flaky.
        from unittest.mock import patch
        with patch('library_app.reports.excel_exporter.ExcelExporter.generate_workbook', side_effect=PermissionError):
            with self.assertRaises(PermissionError) as context:
                self.report_service.generate_report("OWNER", "OwnerUser")
            self.assertIn("Please close TrueEyeball_Library_Report.xlsx", str(context.exception))
            
    def test_staff_restricted_export(self):
        # Staff can export their own
        path, is_new = self.report_service.generate_report("STAFF", "StaffUser")
        self.assertTrue(os.path.exists(path))
        self.assertEqual(os.path.basename(path), "TrueEyeball_Staff_Report.xlsx")
        
        # Staff cannot export FULL report directly via lower level if tampered
        from library_app.permissions import PermissionManager, Permission
        with self.assertRaises(PermissionError):
            PermissionManager.require_permission("STAFF", Permission.EXPORT_FULL_REPORT)

if __name__ == '__main__':
    unittest.main()
