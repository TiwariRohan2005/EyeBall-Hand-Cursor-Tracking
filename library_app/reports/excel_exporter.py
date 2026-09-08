import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import logging
import time

logger = logging.getLogger(__name__)

class ExcelExporter:
    @staticmethod
    def _apply_header_style(sheet, headers):
        sheet.append(headers)
        fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
        font = Font(color="FFFFFF", bold=True)
        for col, _ in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col)
            cell.fill = fill
            cell.font = font
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        
    @staticmethod
    def _format_timestamp(ts):
        if not ts: return ""
        try:
            return time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(float(ts)))
        except:
            return str(ts)

    @classmethod
    def generate_workbook(cls, filepath: str, inventory: list, transactions: list, summary: dict, staff_only: bool = False):
        import os
        
        file_exists = os.path.exists(filepath)
        if file_exists:
            wb = openpyxl.load_workbook(filepath)
        else:
            wb = openpyxl.Workbook()
            
        # Lifecycle Report Sheet
        ws_inv_title = "My Library Activity" if staff_only else "Complete Book Lifecycle"
        
        if ws_inv_title in wb.sheetnames:
            ws_inv = wb[ws_inv_title]
            # Clear existing data rows (keep header)
            if ws_inv.max_row > 1:
                ws_inv.delete_rows(2, ws_inv.max_row - 1)
        else:
            if not file_exists and len(wb.sheetnames) == 1 and wb.sheetnames[0] == 'Sheet':
                ws_inv = wb.active
                ws_inv.title = ws_inv_title
            else:
                ws_inv = wb.create_sheet(ws_inv_title)
            
            headers = ["Book Code", "Title", "Author", "Borrower", "Original Issued By", "Issue Date", "Returned By", "Return Date", "Current Status"]
            if not staff_only:
                headers.insert(3, "Category")
            cls._apply_header_style(ws_inv, headers)
            
        # Group transactions by book_id to reconstruct lifecycle
        book_history = {}
        for t in reversed(transactions): # oldest first inside transactions? they are DESC by timestamp, so reversed() is oldest first.
            bid = t['book_id']
            if bid not in book_history:
                book_history[bid] = {'issue': None, 'return': None}
            if t['action'] == 'ISSUE':
                book_history[bid]['issue'] = t
                book_history[bid]['return'] = None # Clear return if issued again
            elif t['action'] == 'RETURN':
                book_history[bid]['return'] = t

        for b in inventory:
            bid = b['id']
            # If no history exists for this book in this scoped report, and it's staff_only, skip it? 
            # Wait, inventory passed to export_staff is already filtered to `books_involved`
            hist = book_history.get(bid, {})
            iss = hist.get('issue')
            ret = hist.get('return')
            
            borrower = iss.get('user_id') if iss else b.get('issued_to_user_id')
            issued_by = iss.get('staff_user_id') if iss else b.get('issued_by_user_id')
            issue_date = cls._format_timestamp(iss.get('timestamp') if iss else b.get('issue_timestamp'))
            
            returned_by = ret.get('staff_user_id') if ret else None
            return_date = cls._format_timestamp(ret.get('timestamp')) if ret else None
            
            row = [
                b.get('book_code'),
                b.get('title'),
                b.get('author'),
                borrower,
                issued_by,
                issue_date,
                returned_by,
                return_date,
                b.get('status')
            ]
            if not staff_only:
                row.insert(3, b.get('category'))
            ws_inv.append(row)
            
        # Sheet 2: Transactions
        ws_txn_title = "Raw Transactions Log"
        existing_txns = set()
        
        if ws_txn_title in wb.sheetnames:
            ws_txn = wb[ws_txn_title]
            # Build set of existing txn ids from Col A
            for row in ws_txn.iter_rows(min_row=2, max_col=1, values_only=True):
                if row[0] is not None:
                    existing_txns.add(row[0])
        else:
            ws_txn = wb.create_sheet(ws_txn_title)
            cls._apply_header_style(ws_txn, ["ID", "Book ID", "Action", "User ID", "Staff User ID", "Timestamp"])
            
        for t in transactions:
            if t.get('id') not in existing_txns:
                ws_txn.append([
                    t.get('id'),
                    t.get('book_id'),
                    t.get('action'),
                    t.get('user_id'),
                    t.get('staff_user_id'),
                    cls._format_timestamp(t.get('timestamp'))
                ])
                existing_txns.add(t.get('id'))
                
        # Sheet 3: Summary
        ws_sum_title = "Summary Analytics"
        if ws_sum_title in wb.sheetnames:
            ws_sum = wb[ws_sum_title]
            # Clear existing data rows (keep header)
            if ws_sum.max_row > 1:
                ws_sum.delete_rows(2, ws_sum.max_row - 1)
        else:
            ws_sum = wb.create_sheet(ws_sum_title)
            ws_sum.append(["Metric", "Value"])
            ws_sum.cell(row=1, column=1).font = Font(bold=True)
            ws_sum.cell(row=1, column=2).font = Font(bold=True)
            
        for k, v in summary.items():
            ws_sum.append([str(k).replace('_', ' ').title(), v])
            
        # Auto-adjust widths
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter 
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                ws.column_dimensions[column].width = min(40, max_length + 2)
                
        wb.save(filepath)
        logger.info(f"Excel report generated/updated successfully at {filepath}")
        return filepath
